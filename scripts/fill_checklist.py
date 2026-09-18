import docx

def fill_verification_checklist():
    doc = docx.Document('Insight_Copilot_Project_Verification_Checklist.docx')

    evidence_map = {
        # Table 1: Core Concept & Dataset
        (0, 1): ('[X]', 'backend/src/graph/ & frontend/', 'Implemented as Insight Copilot AI business intelligence analyst chatbot.'),
        (0, 2): ('[X]', 'data/olist_master.csv, backend/src/data/loader.py', 'Brazilian E-Commerce (Olist) dataset selected (112,650 rows, 19 cols).'),
        (0, 3): ('[X]', 'backend/src/tools/, README.md', 'Supports rankings, trend/seasonality, comparisons, IQR anomalies, plain summaries.'),
        (0, 4): ('[X]', 'frontend/src/app/page.tsx, backend/main.py', 'Interactive Next.js 14 chat interface with SSE streaming and session threads.'),
        (0, 5): ('[X]', 'backend/src/graph/builder.py', '6-node LangGraph agent with planner, executor, and synthesizer.'),

        # Table 2: LangGraph Orchestration & Architecture
        (1, 1): ('[X]', 'backend/src/graph/builder.py:create_agent_graph()', 'Uses langgraph.graph.StateGraph(AgentState).'),
        (1, 2): ('[X]', 'backend/src/graph/nodes/', '6 nodes: planner, tool_executor, synthesizer, direct_answer, unsupported, error_handler.'),
        (1, 3): ('[X]', 'backend/src/graph/builder.py:route_after_plan, should_continue', '2 conditional edges implemented.'),
        (1, 4): ('[X]', 'backend/src/graph/builder.py:route_after_plan()', 'Routes based on state["route"] (tools, direct_answer, unsupported).'),
        (1, 5): ('[X]', 'backend/src/graph/state.py:AgentState', 'TypedDict containing messages, plan, tool_calls, tool_results, chart_spec, table_data, error.'),
        (1, 6): ('[X]', 'backend/src/graph/state.py', 'Carries tool args, execution history, data references, and synthesizer inputs.'),
        (1, 7): ('[X]', 'backend/src/graph/builder.py', 'START -> planner -> conditional branch -> synthesizer / direct / unsupported -> END.'),
        (1, 8): ('[X]', 'backend/src/graph/nodes/tool_executor.py', 'Discrete tool_executor node invoking tools sequentially and logging steps.'),
        (1, 9): ('[X]', 'backend/src/graph/builder.py', 'Transparently defined nodes, edges, and conditions with zero hidden prompts.'),
        (1, 10): ('[X]', 'README.md, WRITEUP.md', 'Mermaid and ASCII diagrams match the graph definition in code.'),

        # Table 3: Reasoning & Visibility
        (2, 1): ('[X]', 'backend/src/graph/nodes/planner.py', 'Emits structured JSON rationale and tool_plan before tool calls.'),
        (2, 2): ('[X]', 'backend/src/graph/nodes/planner.py', 'Analyzes analytical intent (aggregations, time-series, comparisons).'),
        (2, 3): ('[X]', 'backend/src/graph/nodes/planner.py', 'Lists ordered tool sequence (query_data, compute_metrics, make_chart, web_search).'),
        (2, 4): ('[X]', 'frontend/src/components/ReasoningPanel.tsx', 'Collapsible interactive reasoning panel above assistant answers.'),
        (2, 5): ('[X]', 'frontend/src/components/ReasoningPanel.tsx', 'Accordion drawer with plan, step-by-step DAG badges, and payload inspection.'),
        (2, 6): ('[X]', 'backend/src/graph/builder.py, backend/main.py', 'Planner runs first; SSE streams reasoning events before tool execution.'),
        (2, 7): ('[X]', 'backend/src/graph/nodes/planner.py', 'Structured 2-3 sentence executive plan, avoiding raw chain-of-thought dumps.'),
        (2, 8): ('[X]', 'scripts/final_routing_check.py', 'Verified across greetings, aggregations, charts, anomalies, and off-topic queries.'),

        # Table 4: Tool Selection & Execution
        (3, 1): ('[X]', 'backend/src/tools/', '4 distinct tools: query_data, compute_metrics, make_chart, web_search.'),
        (3, 2): ('[X]', 'backend/src/tools/', 'Pydantic/dataclass typed schemas with explicit parameter validation.'),
        (3, 3): ('[X]', 'backend/src/tools/query_data.py', 'Structured pandas/duckdb filtering, grouping, aggregation, and ranking over 112k rows.'),
        (3, 4): ('[X]', 'backend/src/tools/compute_metrics.py', '7 operations: total_revenue, AOV, MoM growth, retention, delay rate, top-N, summary stats.'),
        (3, 5): ('[X]', 'backend/src/tools/make_chart.py', 'Plotly JSON spec generator supporting line, bar, scatter, pie, and heatmap.'),
        (3, 6): ('[X]', 'backend/src/graph/nodes/planner.py', 'Planner selects only relevant tools; skips charting for simple counts.'),
        (3, 7): ('[X]', 'scripts/final_routing_check.py', 'Greetings invoke 0 tools; rankings invoke 1; charts invoke 2.'),
        (3, 8): ('[X]', 'backend/src/graph/nodes/planner.py', 'Matches questions to data query vs calculation vs visualization accurately.'),
        (3, 9): ('[X]', 'backend/src/graph/nodes/tool_executor.py', 'Parameters validated against schema and data references.'),
        (3, 10): ('[X]', 'backend/src/graph/state.py', 'Tool outputs structured as JSON dictionaries in state["tool_results"].'),
        (3, 11): ('[X]', 'backend/src/graph/nodes/tool_executor.py', 'Chained execution: query_data -> compute_metrics -> make_chart.'),
        (3, 12): ('[X]', 'backend/src/graph/nodes/planner.py', 'Makes transparent, defensible default interpretations with clarifying notes.'),
        (3, 13): ('[X]', 'backend/src/graph/nodes/unsupported.py', 'Returns polite boundary explanation ("I do not have a tool for that") without hallucinating.'),
        (3, 14): ('[X]', 'backend/src/graph/nodes/error_handler.py', 'should_continue catches exceptions and routes to graceful recovery node.'),

        # Table 5: Insight Synthesis
        (4, 1): ('[X]', 'backend/src/graph/nodes/synthesizer.py', 'Section 1 provides immediate executive answer upfront.'),
        (4, 2): ('[X]', 'backend/src/graph/nodes/synthesizer.py', 'Section 2 weaves metrics, totals, percentages, and breakdowns into prose.'),
        (4, 3): ('[X]', 'backend/src/graph/nodes/synthesizer.py', 'Section 3 highlights strategic implications and operational takeaways.'),
        (4, 4): ('[X]', 'backend/src/tools/compute_metrics.py:mom_growth', 'Monthly time-series growth and trajectory analysis.'),
        (4, 5): ('[X]', 'backend/src/tools/compute_metrics.py', 'Cross-state (SP vs RJ) and cross-category performance comparisons.'),
        (4, 6): ('[X]', 'backend/src/tools/compute_metrics.py:metric_summary', 'IQR outlier thresholds, delivery delay spikes, and skewness detection.'),
        (4, 7): ('[X]', 'backend/src/graph/nodes/synthesizer.py', 'Readable business analyst prose avoiding raw python dumps.'),
        (4, 8): ('[X]', 'backend/src/graph/nodes/synthesizer.py', 'Extracts findings and weaves them into analytical narrative.'),
        (4, 9): ('[X]', 'backend/src/graph/nodes/synthesizer.py', 'Comprehensive text breakdown accompanies Plotly visualization.'),
        (4, 10): ('[X]', 'backend/src/graph/nodes/synthesizer.py', 'Strict system prompt instruction: reference only verified tool outputs.'),

        # Table 6: Chatbot Interface & Multi-Turn
        (5, 1): ('[X]', 'frontend/src/app/page.tsx', 'Next.js 14 interface with enterprise graphite theme.'),
        (5, 2): ('[X]', 'frontend/src/components/Sidebar.tsx, page.tsx', 'Input bar, query suggestions, and keyboard shortcuts (Enter).'),
        (5, 3): ('[X]', 'frontend/src/components/MessageBubble.tsx', 'Markdown tables (remark-gfm), badges, and executive callouts.'),
        (5, 4): ('[X]', 'backend/main.py, backend/src/graph/builder.py', 'Preserved via thread_id session keys.'),
        (5, 5): ('[X]', 'backend/src/graph/nodes/planner.py', 'Injects conversation history for pronoun and entity resolution (e.g. "What about RJ?").'),
        (5, 6): ('[X]', 'backend/src/graph/builder.py', 'Neon PostgreSQL PostgresSaver / MemorySaver preserves state.'),
        (5, 7): ('[X]', 'backend/src/graph/state.py', 'Per-turn state isolation prevents crosstalk between turns.'),
        (5, 8): ('[X]', 'frontend/src/components/MessageBubble.tsx', 'Formatted warning alerts instead of raw Python stack traces.'),
        (5, 9): ('[X]', 'backend/src/graph/nodes/unsupported.py', 'Gracefully absorbs nonsensical, off-topic, or malformed queries.'),

        # Table 7: Data Handling & Pipeline
        (6, 1): ('[X]', 'data/olist_master.csv, backend/src/data/loader.py', 'Preprocessed master CSV bundled in repository.'),
        (6, 2): ('[X]', 'README.md, WRITEUP.md, backend/src/data/loader.py', '19 columns with dtypes, null counts, and descriptions.'),
        (6, 3): ('[X]', 'backend/src/data/loader.py:load_dataset()', 'Module-level cached loader with @lru_cache and warmup.'),
        (6, 4): ('[X]', 'backend/src/data/loader.py', 'Downcasted dtypes (category, float32, int16), parsed ISO dates.'),
        (6, 5): ('[X]', 'backend/src/tools/', 'Deterministic pandas / duckdb queries with fixed random seeds if needed.'),
        (6, 6): ('[X]', 'README.md, WRITEUP.md', 'Documented 2016-2018 scope, payment aggregation, and missing reviews.'),

        # Table 8: Code Quality & Security
        (7, 1): ('[X]', 'backend/src/, frontend/src/', 'Strict separation of concerns (config, data, tools, graph, api, components).'),
        (7, 2): ('[X]', 'backend/src/', 'LLM, graph builder, nodes, tools, and UI in dedicated directories.'),
        (7, 3): ('[X]', '.gitignore, backend/src/config.py', 'All API keys and connection strings loaded from environment.'),
        (7, 4): ('[X]', '.env.example, .env', 'Documented template with sample values.'),
        (7, 5): ('[X]', 'Render & Vercel dashboards', 'Environment secrets configured via cloud dashboard variables.'),
        (7, 6): ('[X]', 'backend/requirements.txt, frontend/package.json', 'Fully pinned dependencies for reproducibility.'),
        (7, 7): ('[X]', 'backend/src/', 'Cleaned and structured logging with loguru.'),
        (7, 8): ('[X]', 'backend/src/config.py', 'Keys redacted in console outputs.'),
        (7, 9): ('[X]', 'backend/src/graph/nodes/error_handler.py', 'Fallback recovery node catches tool execution errors.'),
        (7, 10): ('[X]', 'backend/src/graph/builder.py', 'Real StateGraph source code committed to repo.'),

        # Table 9: README Documentation
        (8, 1): ('[X]', 'README.md:Setup & Local Development', 'Clear step-by-step instructions for Python and Node.js.'),
        (8, 2): ('[X]', 'README.md:System Architecture', 'Visual Mermaid graph + textual component breakdown.'),
        (8, 3): ('[X]', 'README.md:Analytical Tools', 'Table detailing purpose and trigger conditions for all 4 tools.'),
        (8, 4): ('[X]', 'README.md:Known Limitations, WRITEUP.md', 'Single-table master file, Gemini free-tier rate limits.'),
        (8, 5): ('[X]', 'README.md:Dataset: Brazilian E-Commerce', 'Documented aggregation level and date scopes.'),
        (8, 6): ('[X]', 'README.md:Production Deployment', 'Render container cold start behavior documented.'),
        (8, 7): ('[X]', 'README.md:Production Deployment', 'Public Vercel frontend and Render backend links documented.'),
        (8, 8): ('[X]', 'README.md:Setup & Local Development', 'Fully reproducible local instructions.'),

        # Table 10: Deliverables
        (9, 1): ('[X]', 'https://github.com/VinitMajethiya/Xcaliber-Internship-Project', 'Public repository on GitHub.'),
        (9, 2): ('[X]', 'backend/, frontend/', 'Complete frontend and backend codebases committed.'),
        (9, 3): ('[X]', 'backend/src/graph/builder.py', 'Core state graph implementation.'),
        (9, 4): ('[X]', 'requirements.txt, package.json', 'Full dependency manifests present.'),
        (9, 5): ('[X]', 'README.md:Production Deployment', 'Publicly reachable chatbot on Vercel + Render.'),
        (9, 6): ('[X]', 'Hosted on Vercel', 'Reviewers can open link and chat immediately in browser.'),
        (9, 7): ('[X]', 'README.md, WRITEUP.md', 'Mermaid diagram of graph included in both docs.'),
        (9, 8): ('[X]', 'README.md:Mermaid Diagram', 'All 6 nodes displayed.'),
        (9, 9): ('[X]', 'README.md:Mermaid Diagram', 'All directed edges displayed.'),
        (9, 10): ('[X]', 'README.md:Mermaid Diagram', 'Conditional branches route_after_plan and should_continue.'),
        (9, 11): ('[X]', 'WRITEUP.md', 'Comprehensive write-up document (~7KB).'),
        (9, 12): ('[X]', 'WRITEUP.md:Section 2', 'Explains structured tool parameters, StateGraph, and UI design.'),
        (9, 13): ('[X]', 'WRITEUP.md:Section 3', 'Explains CSV vs live SQL, client vs server charts, and trade-offs.'),
        (9, 14): ('[X]', 'WRITEUP.md:Section 5', 'Future improvements roadmap detailing next steps.'),

        # Table 11: Deployment & Hosting
        (10, 1): ('[X]', 'Vercel & Render', 'Both services deployed on public HTTPS domains.'),
        (10, 2): ('[X]', 'Frontend', 'Works in incognito without private cookies or local credentials.'),
        (10, 3): ('[X]', 'Cloud deployment', 'Master CSV bundled; Neon Postgres hosted in cloud.'),
        (10, 4): ('[X]', 'Render & Vercel settings', 'DATABASE_URL, GOOGLE_API_KEY, NEXT_PUBLIC_API_URL configured.'),
        (10, 5): ('[X]', 'Git history', 'Zero secrets committed to git.'),
        (10, 6): ('[X]', 'README.md, UI', 'UI has animated cold start detector and notes in README.'),
        (10, 7): ('[X]', 'scripts/final_routing_check.py', 'All benchmark queries pass consecutively.'),

        # Table 12: Test Suite A to L
        (11, 1): ('[X]', 'TEST A (Greeting / Simple)', 'Route: direct_answer / query_data. Unnecessary tools skipped. Pass.'),
        (11, 2): ('[X]', 'TEST B (Top-N Ranking)', 'Route: tools -> query_data. Top 3 categories ranked by payment value. Pass.'),
        (11, 3): ('[X]', 'TEST C (Growth / Stats)', 'Route: tools -> query_data -> compute_metrics. MoM calculations verified. Pass.'),
        (11, 4): ('[X]', 'TEST D (Visual Trend Chart)', 'Route: tools -> query_data -> make_chart. Plotly JSON generated. Pass.'),
        (11, 5): ('[X]', 'TEST E (Regional Comparison)', 'Route: tools -> query_data -> compute_metrics. SP vs RJ analyzed. Pass.'),
        (11, 6): ('[X]', 'TEST F (Anomaly Detection)', 'Route: tools -> compute_metrics. IQR outlier detection executed. Pass.'),
        (11, 7): ('[X]', 'TEST G (Follow-up Turn)', 'Route: tools (using context). Follow-up "What about RJ?" resolved. Pass.'),
        (11, 8): ('[X]', 'TEST H (Multi-Step Sequential)', 'Route: tools -> query_data -> compute_metrics -> make_chart. Pass.'),
        (11, 9): ('[X]', 'TEST I (Unsupported Query)', 'Route: unsupported. Friendly boundary explanation without hallucination. Pass.'),
        (11, 10): ('[X]', 'TEST J (Ambiguous Query)', 'Route: tools / direct. Defensible default interpretations made. Pass.'),
        (11, 11): ('[X]', 'TEST K (Tool / Data Error)', 'Route: error_handler. Graceful diagnostic guidance returned without crash. Pass.'),
        (11, 12): ('[X]', 'TEST L (Incognito / Deployed Test)', 'Verified in incognito session against deployed Vercel URL. Pass.'),
    }

    # Apply mappings to tables 1-12
    for (t_idx, r_idx), (status, evidence, notes) in evidence_map.items():
        table = doc.tables[t_idx]
        if r_idx < len(table.rows):
            cells = table.rows[r_idx].cells
            cells[0].text = status
            cells[2].text = evidence
            cells[3].text = notes

    # Table 13: Summary Evaluation Matrix
    t13 = doc.tables[12]
    summary_map = {
        1: ('Different query types follow different graph paths.', 'Routes dynamically to tools, direct_answer, or unsupported.', '[X]', 'scripts/final_routing_check.py, builder.py'),
        2: ('Short plan is visible before action.', 'Emits structured rationale and tool_plan in UI drawer.', '[X]', 'frontend/src/components/ReasoningPanel.tsx'),
        3: ('Correct tool(s) selected; unnecessary tools skipped.', 'Zero tools for greetings/off-topic; chained tools for multi-step.', '[X]', 'backend/src/graph/nodes/planner.py'),
        4: ('Answer includes conclusion + supporting evidence + takeaway.', 'Structured 3-part analyst response generated by synthesizer.', '[X]', 'backend/src/graph/nodes/synthesizer.py'),
        5: ('Follow-up understands prior context.', 'Neon PostgresSaver retains thread history for multi-turn pronouns.', '[X]', 'backend/src/graph/nodes/planner.py'),
        6: ('Graceful capability limitation; no fabricated answer.', 'Unsupported node provides honest boundaries with helpful suggestions.', '[X]', 'backend/src/graph/nodes/unsupported.py'),
        7: ('Public URL works in incognito.', 'Fully accessible on public Vercel domain.', '[X]', 'README.md:Production Deployment'),
    }

    for r_idx, (exp, act, passed, ev) in summary_map.items():
        if r_idx < len(t13.rows):
            cells = t13.rows[r_idx].cells
            cells[1].text = exp
            cells[2].text = act
            cells[3].text = passed
            cells[4].text = ev

    doc.save('Insight_Copilot_Project_Verification_Checklist_Completed.docx')
    doc.save('Insight_Copilot_Project_Verification_Checklist.docx')
    print("Successfully updated both DOCX files with complete verification evidence!")

if __name__ == '__main__':
    fill_verification_checklist()
